from __future__ import annotations

import argparse
import io
import re
import sqlite3
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import openpyxl

try:
    from Backend.db import DEFAULT_DB_PATH, connect, init_db, parse_entry_registration, resolve_study_program_id
    from Backend.excel_reader import (
        clean_demographic_value,
        clean_excel_text,
        normalize_imported_name,
        normalize_name,
        normalize_nim,
        normalize_text,
        read_sheet,
        read_sheet_headers,
        workbook_sheet_names,
    )
except ModuleNotFoundError:
    from db import DEFAULT_DB_PATH, connect, init_db, parse_entry_registration, resolve_study_program_id
    from excel_reader import (
        clean_demographic_value,
        clean_excel_text,
        normalize_imported_name,
        normalize_name,
        normalize_nim,
        normalize_text,
        read_sheet,
        read_sheet_headers,
        workbook_sheet_names,
    )

DEFAULT_WORKBOOK = Path(__file__).resolve().parents[1] / "MASTER_DATA_2023_1_2026_1.xlsx"
DEFAULT_PERIOD = "UKT 2023.1 s/d 2025.2"
DEFAULT_BILL_TYPE = "UKT BRIVA"
DEFAULT_CURRENT_PERIOD = "Semester Ganjil 2026"
DEFAULT_INSTRUCTIONS = (
    "Bayar melalui BRIVA BRI dengan nomor VA yang tampil. "
    "Pastikan nama dan nominal sesuai sebelum menyelesaikan pembayaran."
)
REQUIRED_SYNC_HEADERS = ("NIM", "Nama Mahasiswa", "BRIVA", "Jumlah")
OPTIONAL_ISSUE_HEADERS = ("NIM", "Nama Mahasiswa", "BRIVA", "Jumlah", "Keterangan")
CURRENT_REQUIRED_HEADERS = ("NIM", "Nama", "No Rek", "Jumlah")

MASTER_TEMPLATE_HEADERS = [
    "NIM",
    "Nama",
    "NO KTP",
    "Tempat Lahir",
    "Tanggal Lahir",
    "Nama Ibu Kandung",
    "e-Mail",
    "No Kontak",
    "Registrasi Awal",
    "Program Studi",
    "No Rek",
    "Jumlah",
    "Batas Pembayaran",
]

MASTER_SAMPLE_ROWS = [
    [
        "049530265",
        "Muhamad Romli",
        "3603100510860014",
        "Tangerang",
        "14 September 2000",
        "Siti Aminah",
        "rhomly0496@gmail.com",
        "082310867195",
        "UNIVERSITAS TERBUKA 2023.1",
        "FEB - Akuntansi",
        "178100023200085",
        1850000,
        "22 Januari 2027 Pukul 11.59 WIB",
    ],
    [
        "049532688",
        "Ria Anggraeni",
        "3603115601060002",
        "Serang",
        "25 Mei 2001",
        "Nurjanah",
        "riaa1390@gmail.com",
        "0895411921596",
        "UNIVERSITAS TERBUKA 2023.2",
        "FHISIP - Sosiologi",
        "178100023200060",
        1850000,
        "22 Januari 2027 Pukul 11.59 WIB",
    ],
]


def generate_master_data_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_Data_Mahasiswa"
    ws.append(MASTER_TEMPLATE_HEADERS)
    for sample in MASTER_SAMPLE_ROWS:
        ws.append(sample)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass(frozen=True)
class ImportLayout:
    kind: str
    data_sheet: str
    issue_sheet: str | None
    headers: dict[str, str]
    default_period: str


def _amount_to_int(value: str) -> int | None:
    cleaned = re.sub(r"\D+", "", clean_excel_text(value))
    if not cleaned or not cleaned.isdigit():
        return None
    return int(cleaned)


def _normalized_headers(headers: list[str]) -> dict[str, str]:
    return {
        normalize_text(header).casefold(): normalize_text(header)
        for header in headers
        if normalize_text(header)
    }


def _find_header_key(headers: dict[str, str], aliases: tuple[str, ...]) -> str:
    for alias in aliases:
        key = alias.casefold()
        if key in headers:
            return headers[key]
    return ""


def _require_headers(workbook: Path) -> ImportLayout:
    sheet_names = workbook_sheet_names(workbook)
    if "Data Sinkron" in sheet_names and "Data Belum Lengkap" in sheet_names:
        sync_headers = _normalized_headers(read_sheet_headers(workbook, "Data Sinkron"))
        missing = [header for header in REQUIRED_SYNC_HEADERS if header.casefold() not in sync_headers]
        if missing:
            raise ValueError(
                "Struktur header sheet Data Sinkron tidak sesuai. "
                f"Kolom wajib: {', '.join(REQUIRED_SYNC_HEADERS)}. "
                f"Kolom belum ditemukan: {', '.join(missing)}."
            )

        issue_headers = _normalized_headers(read_sheet_headers(workbook, "Data Belum Lengkap"))
        missing_issue_headers = [header for header in OPTIONAL_ISSUE_HEADERS if header.casefold() not in issue_headers]
        if missing_issue_headers:
            raise ValueError(
                "Struktur header sheet Data Belum Lengkap tidak sesuai. "
                f"Kolom wajib: {', '.join(OPTIONAL_ISSUE_HEADERS)}. "
                f"Kolom belum ditemukan: {', '.join(missing_issue_headers)}."
            )
        return ImportLayout("legacy", "Data Sinkron", "Data Belum Lengkap", sync_headers, DEFAULT_PERIOD)

    for sheet_name in sheet_names:
        headers = _normalized_headers(read_sheet_headers(workbook, sheet_name))
        has_nim = any(a in headers for a in ("nim", "nomor induk mahasiswa"))
        has_name = any(a in headers for a in ("nama", "nama mahasiswa", "nama lengkap"))
        has_briva = any(a in headers for a in ("no rek", "no. rek", "no rekening", "briva", "nomor briva", "va"))
        has_amount = any(a in headers for a in ("jumlah", "nominal", "tagihan", "biaya"))
        if has_nim and has_name and has_briva and has_amount:
            return ImportLayout("current", sheet_name, None, headers, DEFAULT_CURRENT_PERIOD)

    raise ValueError(
        "Struktur Excel tidak dikenali. Gunakan format Data Sinkron lama atau format Master Data "
        "dengan kolom NIM, Nama, No Rek, dan Jumlah."
    )


def _get_existing_id(conn: sqlite3.Connection, table: str, column: str, value: str) -> str | None:
    row = conn.execute(f"select id from {table} where {column} = ?", (value,)).fetchone()
    return str(row["id"]) if row else None


def _record_value(record: dict[str, str], headers: dict[str, str], aliases: str | tuple[str, ...]) -> str:
    if isinstance(aliases, str):
        aliases = (aliases,)
    for alias in aliases:
        key = alias.casefold()
        if key in headers:
            return record.get(headers[key], "")
    return ""


def _normalize_briva(value: object) -> str:
    return normalize_nim(value)


def _read_sync_rows(
    workbook: Path,
    layout: ImportLayout,
    period: str,
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], int, list[dict[str, object]]]:
    rows: list[dict[str, object]] = []
    errors: list[dict[str, object]] = []
    sample: list[dict[str, object]] = []
    identity_conflict_rows = 0
    skipped_issues: list[dict[str, object]] = []

    for record in read_sheet(workbook, layout.data_sheet):
        nim = normalize_nim(_record_value(record, layout.headers, ("NIM", "Nomor Induk Mahasiswa")))
        full_name = normalize_imported_name(_record_value(record, layout.headers, ("Nama", "Nama Mahasiswa", "Nama Lengkap")))
        briva = _normalize_briva(_record_value(record, layout.headers, ("No Rek", "No. Rek", "No Rekening", "BRIVA", "Nomor BRIVA", "VA")))
        amount = _amount_to_int(_record_value(record, layout.headers, ("Jumlah", "Nominal", "Tagihan", "Biaya")))
        row_number = int(record.get("_row_number") or 0)
        if not nim or not full_name or not briva or amount is None:
            message = "Baris dilewati karena NIM, nama, BRIVA, atau nominal tidak valid."
            errors.append(
                {
                    "sheet": layout.data_sheet,
                    "row_number": row_number,
                    "severity": "warning",
                    "message": message,
                }
            )
            skipped_issues.append(
                {
                    "sheet_name": layout.data_sheet,
                    "row_number": row_number,
                    "nim": nim,
                    "full_name": full_name,
                    "briva": briva,
                    "amount": clean_excel_text(_record_value(record, layout.headers, ("Jumlah", "Nominal"))),
                    "note": message,
                }
            )
            continue

        raw_ktp = _record_value(record, layout.headers, ("NO KTP", "No. KTP", "NIK", "KTP", "Nomor KTP"))
        raw_tempat = _record_value(record, layout.headers, ("Tempat Lahir", "Tempat_Lahir"))
        raw_tgl = _record_value(record, layout.headers, ("Tanggal Lahir", "Tanggal_Lahir", "Tgl Lahir"))
        raw_ibu = _record_value(record, layout.headers, ("Nama Ibu Kandung", "Nama Ibu", "Ibu Kandung"))
        raw_email = _record_value(record, layout.headers, ("e-Mail", "Email", "E-mail", "Surel"))
        raw_kontak = _record_value(record, layout.headers, ("No Kontak", "No. Kontak", "No Hp", "No. HP", "No Telepon", "Telepon"))
        raw_reg = _record_value(record, layout.headers, ("Registrasi Awal", "Periode Masuk", "Registrasi_Awal"))
        raw_prodi = _record_value(record, layout.headers, ("Program Studi", "Prodi", "Jurusan"))
        raw_due = _record_value(record, layout.headers, ("Batas Pembayaran", "Jatuh Tempo", "Due Date"))

        initial_reg = clean_demographic_value(raw_reg)
        entry_year, entry_semester, entry_period = parse_entry_registration(initial_reg)

        row = {
            "nim": nim,
            "full_name": full_name,
            "briva": briva,
            "amount": amount,
            "row_number": row_number,
            "period": period,
            "no_ktp": clean_demographic_value(raw_ktp),
            "tempat_lahir": clean_demographic_value(raw_tempat),
            "tanggal_lahir": clean_demographic_value(raw_tgl),
            "nama_ibu_kandung": clean_demographic_value(raw_ibu),
            "email": clean_demographic_value(raw_email),
            "phone_number": normalize_nim(raw_kontak) if raw_kontak and clean_demographic_value(raw_kontak) else None,
            "initial_registration": initial_reg,
            "entry_year": entry_year,
            "entry_semester": entry_semester,
            "entry_period": entry_period,
            "program_study": clean_demographic_value(raw_prodi),
            "due_date": clean_demographic_value(raw_due),
        }
        rows.append(row)
        if len(sample) < 5:
            sample.append({key: row[key] for key in ("nim", "full_name", "briva", "amount", "program_study", "due_date")})

    canonical_profiles: dict[str, dict[str, object]] = {}
    for row in rows:
        nim = str(row["nim"])
        canonical = canonical_profiles.get(nim)
        if canonical is None:
            canonical_profiles[nim] = row
            continue
        if normalize_name(str(canonical["full_name"])) != normalize_name(str(row["full_name"])):
            identity_conflict_rows += 1
            errors.append(
                {
                    "sheet": layout.data_sheet,
                    "row_number": row["row_number"],
                    "severity": "warning",
                    "message": "NIM muncul dengan nama berbeda. Nama pada baris pertama digunakan untuk profil mahasiswa.",
                }
            )
        for field in (
            "full_name", "no_ktp", "tempat_lahir", "tanggal_lahir", "nama_ibu_kandung",
            "email", "phone_number", "program_study", "initial_registration",
            "entry_year", "entry_semester", "entry_period",
        ):
            if canonical.get(field):
                row[field] = canonical[field]

    return rows, errors, sample, identity_conflict_rows, skipped_issues


def _existing_bills(
    db_path: str | Path,
    nims: set[str],
    period: str,
    source_file: str,
) -> tuple[dict[str, list[sqlite3.Row]], dict[str, list[sqlite3.Row]], dict[int, sqlite3.Row]]:
    if not nims:
        return {}, {}, {}

    conn = connect(db_path)
    init_db(conn)
    placeholders = ",".join("?" for _ in nims)
    rows = conn.execute(
        f"""
        select b.id, b.student_id, b.briva, b.amount, b.period, b.status, b.source_file, b.source_row_number,
               b.due_date, s.nim, s.full_name, s.program_study, s.initial_registration, s.phone_number
        from bills b
        join students s on s.id = b.student_id
        where (b.source_file = ? or b.period = ? or s.nim in ({placeholders}))
          and b.deleted_at is null
          and s.deleted_at is null
        """,
        (source_file, period, *sorted(nims)),
    ).fetchall()
    conn.close()

    by_briva: dict[str, list[sqlite3.Row]] = {}
    by_nim: dict[str, list[sqlite3.Row]] = {}
    by_source_row: dict[int, sqlite3.Row] = {}
    for row in rows:
        by_briva.setdefault(str(row["briva"]), []).append(row)
        if row["period"] == period:
            by_nim.setdefault(str(row["nim"]), []).append(row)
        if row["source_file"] == source_file and row["source_row_number"] is not None:
            by_source_row[int(row["source_row_number"])] = row
    return by_briva, by_nim, by_source_row


def _append_limited(items: list[dict[str, object]], item: dict[str, object], limit: int = 12) -> None:
    if len(items) < limit:
        items.append(item)


def _analyze_workbook(
    workbook_path: str | Path,
    db_path: str | Path | None = None,
    period: str | None = None,
    source_file_name: str | None = None,
) -> dict[str, object]:
    workbook = Path(workbook_path)
    if not workbook.exists():
        raise FileNotFoundError(f"File Excel tidak ditemukan: {workbook}")

    layout = _require_headers(workbook)
    effective_period = period or layout.default_period
    source_file = source_file_name or workbook.name
    rows, errors, sample, identity_conflict_rows, skipped_issues = _read_sync_rows(workbook, layout, effective_period)
    valid_rows = len(rows)
    issue_rows = len(errors)
    critical_rows = sum(1 for error in errors if error["severity"] == "critical")
    briva_counts = Counter(str(row["briva"]) for row in rows)
    nim_counts = Counter(str(row["nim"]) for row in rows)
    duplicate_briva_conflict_rows = 0
    multiple_bill_rows = 0

    for row in rows:
        rows_with_same_briva = [candidate for candidate in rows if candidate["briva"] == row["briva"]]
        nims_for_same_briva = {str(candidate["nim"]) for candidate in rows_with_same_briva}
        if briva_counts[str(row["briva"])] > 1 and len(nims_for_same_briva) > 1:
            duplicate_briva_conflict_rows += 1
            critical_rows += 1
            issue_rows += 1
            _append_limited(
                errors,
                {
                    "sheet": layout.data_sheet,
                    "row_number": row["row_number"],
                    "severity": "critical",
                    "message": "BRIVA yang sama muncul untuk NIM berbeda dalam file.",
                },
            )
        if nim_counts[str(row["nim"])] > 1:
            multiple_bill_rows += 1
            issue_rows += 1
            _append_limited(
                errors,
                {
                    "sheet": layout.data_sheet,
                    "row_number": row["row_number"],
                    "severity": "warning",
                    "message": "NIM muncul lebih dari satu kali. Sistem akan menyimpan sebagai beberapa tagihan.",
                },
            )

    by_briva: dict[str, list[sqlite3.Row]] = {}
    by_nim: dict[str, list[sqlite3.Row]] = {}
    by_source_row: dict[int, sqlite3.Row] = {}
    if db_path is not None:
        by_briva, by_nim, by_source_row = _existing_bills(
            db_path, {str(row["nim"]) for row in rows}, effective_period, source_file
        )

    new_rows = 0
    unchanged_rows = 0
    update_rows = 0
    amount_change_rows = 0
    briva_change_rows = 0
    conflict_rows = 0
    requires_update_confirmation = False
    changes: list[dict[str, object]] = []
    actions: list[dict[str, object]] = []
    used_existing_bill_ids: set[str] = set()

    for row in rows:
        nim = str(row["nim"])
        briva = str(row["briva"])
        amount = int(row["amount"])
        row_number = int(row["row_number"])
        existing_source_row = by_source_row.get(row_number)
        matching_briva_rows = [
            candidate
            for candidate in by_briva.get(briva, [])
            if candidate["nim"] == nim and candidate["period"] == effective_period and str(candidate["id"]) not in used_existing_bill_ids
        ]
        conflicting_briva_rows = [candidate for candidate in by_briva.get(briva, []) if candidate["nim"] != nim]
        existing_briva = matching_briva_rows[0] if matching_briva_rows else None
        current_bills = [candidate for candidate in by_nim.get(nim, []) if str(candidate["id"]) not in used_existing_bill_ids]

        if existing_source_row:
            if existing_source_row["nim"] != nim:
                critical_rows += 1
                issue_rows += 1
                conflict_rows += 1
                _append_limited(
                    errors,
                    {
                        "sheet": layout.data_sheet,
                        "row_number": row["row_number"],
                        "severity": "critical",
                        "message": "Baris sumber file ini sebelumnya terdaftar untuk NIM lain.",
                    },
                )
                continue
            if str(existing_source_row["id"]) not in used_existing_bill_ids:
                existing_briva = existing_source_row

        if conflicting_briva_rows:
            critical_rows += 1
            issue_rows += 1
            conflict_rows += 1
            _append_limited(
                errors,
                {
                    "sheet": layout.data_sheet,
                    "row_number": row["row_number"],
                    "severity": "critical",
                    "message": "BRIVA sudah terdaftar untuk NIM lain.",
                },
            )
            continue

        if existing_briva:
            if existing_briva["period"] != effective_period:
                critical_rows += 1
                issue_rows += 1
                conflict_rows += 1
                _append_limited(
                    errors,
                    {
                        "sheet": layout.data_sheet,
                        "row_number": row["row_number"],
                        "severity": "critical",
                        "message": "BRIVA sudah terdaftar pada periode lain.",
                    },
                )
                continue

            amount_changed = int(existing_briva["amount"]) != amount
            name_changed = normalize_name(str(existing_briva["full_name"])) != normalize_name(str(row["full_name"]))
            program_changed = normalize_text(existing_briva["program_study"]) != normalize_text(row["program_study"])
            registration_changed = normalize_text(existing_briva["initial_registration"]) != normalize_text(row["initial_registration"])
            phone_changed = normalize_text(existing_briva["phone_number"]) != normalize_text(row["phone_number"])
            due_date_changed = normalize_text(existing_briva["due_date"]) != normalize_text(row["due_date"])
            if existing_briva["status"] != "unpaid" and amount_changed:
                critical_rows += 1
                issue_rows += 1
                conflict_rows += 1
                _append_limited(
                    errors,
                    {
                        "sheet": layout.data_sheet,
                        "row_number": row["row_number"],
                        "severity": "critical",
                        "message": "Nominal tagihan yang sudah lunas atau dicicil tidak boleh diubah melalui import.",
                    },
                )
                continue

            if not any((amount_changed, name_changed, program_changed, registration_changed, phone_changed, due_date_changed)):
                unchanged_rows += 1
                actions.append({"type": "unchanged", "row": row, "existing": existing_briva})
                used_existing_bill_ids.add(str(existing_briva["id"]))
                continue

            update_rows += 1
            if amount_changed:
                amount_change_rows += 1
                requires_update_confirmation = True
            action = "update_amount" if amount_changed else "update_profile"
            actions.append({"type": action, "row": row, "existing": existing_briva})
            used_existing_bill_ids.add(str(existing_briva["id"]))
            if len(changes) < 10:
                changes.append(
                    {
                        "row_number": row["row_number"],
                        "nim": nim,
                        "full_name": row["full_name"],
                        "change_type": "nominal" if amount_changed else "nama",
                        "old_briva": existing_briva["briva"],
                        "new_briva": briva,
                        "old_amount": existing_briva["amount"],
                        "new_amount": amount,
                    }
                )
            continue

        has_multiple_rows_for_nim = nim_counts[nim] > 1
        if len(current_bills) == 0 or has_multiple_rows_for_nim:
            new_rows += 1
            actions.append({"type": "new", "row": row, "existing": None})
            continue

        if len(current_bills) > 1:
            new_rows += 1
            actions.append({"type": "new", "row": row, "existing": None})
            continue

        existing = current_bills[0]
        if existing["status"] != "unpaid":
            critical_rows += 1
            issue_rows += 1
            conflict_rows += 1
            _append_limited(
                errors,
                {
                    "sheet": layout.data_sheet,
                    "row_number": row["row_number"],
                    "severity": "critical",
                    "message": "BRIVA tagihan yang sudah lunas atau dicicil tidak boleh diganti melalui import.",
                },
            )
            continue

        update_rows += 1
        briva_change_rows += 1
        requires_update_confirmation = True
        if int(existing["amount"]) != amount:
            amount_change_rows += 1
        actions.append({"type": "replace_briva", "row": row, "existing": existing})
        used_existing_bill_ids.add(str(existing["id"]))
        if len(changes) < 10:
            changes.append(
                {
                    "row_number": row["row_number"],
                    "nim": nim,
                    "full_name": row["full_name"],
                    "change_type": "BRIVA dan nominal" if int(existing["amount"]) != amount else "BRIVA",
                    "old_briva": existing["briva"],
                    "new_briva": briva,
                    "old_amount": existing["amount"],
                    "new_amount": amount,
                }
            )

    if layout.issue_sheet:
        for record in read_sheet(workbook, layout.issue_sheet):
            issue_rows += 1
            _append_limited(
                errors,
                {
                    "sheet": layout.issue_sheet,
                    "row_number": int(record.get("_row_number") or 0),
                    "severity": "warning",
                    "message": normalize_text(record.get("Keterangan")) or "Data belum lengkap.",
                },
            )

    return {
        "valid_rows": valid_rows,
        "critical_rows": critical_rows,
        "issue_rows": issue_rows,
        "new_rows": new_rows,
        "unchanged_rows": unchanged_rows,
        "update_rows": update_rows,
        "amount_change_rows": amount_change_rows,
        "briva_change_rows": briva_change_rows,
        "duplicate_briva_conflict_rows": duplicate_briva_conflict_rows,
        "multiple_bill_rows": multiple_bill_rows,
        "identity_conflict_rows": identity_conflict_rows,
        "skipped_rows": len(skipped_issues),
        "conflict_rows": conflict_rows,
        "requires_update_confirmation": requires_update_confirmation,
        "sample": sample,
        "changes": changes,
        "errors": errors,
        "_skipped_issues": skipped_issues,
        "actions": actions,
    }


def preview_workbook(
    workbook_path: str | Path,
    db_path: str | Path | None = None,
    period: str | None = None,
    source_file_name: str | None = None,
) -> dict[str, object]:
    analysis = _analyze_workbook(workbook_path, db_path, period, source_file_name)
    return {key: value for key, value in analysis.items() if key not in {"actions", "_skipped_issues"}}


def _upsert_student(conn: sqlite3.Connection, row: dict[str, object]) -> str:
    nim = str(row["nim"])
    full_name = normalize_imported_name(row["full_name"])
    student_id = _get_existing_id(conn, "students", "nim", nim) or str(uuid.uuid4())

    no_ktp = clean_demographic_value(row.get("no_ktp"))
    tempat_lahir = clean_demographic_value(row.get("tempat_lahir"))
    tanggal_lahir = clean_demographic_value(row.get("tanggal_lahir"))
    nama_ibu_kandung = clean_demographic_value(row.get("nama_ibu_kandung"))
    email = clean_demographic_value(row.get("email"))
    phone_number = normalize_nim(row.get("phone_number")) if row.get("phone_number") else None
    initial_reg = clean_demographic_value(row.get("initial_registration"))
    program_study = clean_demographic_value(row.get("program_study"))

    entry_year = row.get("entry_year")
    entry_semester = row.get("entry_semester")
    entry_period = row.get("entry_period")
    if entry_year is None or entry_semester is None or entry_period is None:
        y, s, p = parse_entry_registration(initial_reg)
        entry_year = entry_year or y
        entry_semester = entry_semester or s
        entry_period = entry_period or p

    conn.execute(
        """
        insert into students (
            id, nim, full_name, name_norm, no_ktp, tempat_lahir, tanggal_lahir, nama_ibu_kandung,
            program_study, initial_registration, entry_year, entry_semester, entry_period,
            phone_number, email, academic_status, updated_at
        )
        values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'aktif', datetime('now'))
        on conflict(nim) do update set
          full_name = excluded.full_name,
          name_norm = excluded.name_norm,
          no_ktp = case when excluded.no_ktp is not null and excluded.no_ktp <> '' then excluded.no_ktp else students.no_ktp end,
          tempat_lahir = case when excluded.tempat_lahir is not null and excluded.tempat_lahir <> '' then excluded.tempat_lahir else students.tempat_lahir end,
          tanggal_lahir = case when excluded.tanggal_lahir is not null and excluded.tanggal_lahir <> '' then excluded.tanggal_lahir else students.tanggal_lahir end,
          nama_ibu_kandung = case when excluded.nama_ibu_kandung is not null and excluded.nama_ibu_kandung <> '' then excluded.nama_ibu_kandung else students.nama_ibu_kandung end,
          program_study = case when excluded.program_study is not null and excluded.program_study <> '' then excluded.program_study else students.program_study end,
          initial_registration = case when excluded.initial_registration is not null and excluded.initial_registration <> '' then excluded.initial_registration else students.initial_registration end,
          entry_year = case when excluded.entry_year is not null then excluded.entry_year else students.entry_year end,
          entry_semester = case when excluded.entry_semester is not null then excluded.entry_semester else students.entry_semester end,
          entry_period = case when excluded.entry_period is not null then excluded.entry_period else students.entry_period end,
          phone_number = case when excluded.phone_number is not null and excluded.phone_number <> '' then excluded.phone_number else students.phone_number end,
          email = case when excluded.email is not null and excluded.email <> '' then excluded.email else students.email end,
          updated_at = datetime('now')
        """,
        (
            student_id,
            nim,
            full_name,
            normalize_name(full_name),
            no_ktp,
            tempat_lahir,
            tanggal_lahir,
            nama_ibu_kandung,
            program_study,
            initial_reg,
            entry_year,
            entry_semester,
            entry_period,
            phone_number,
            email,
        ),
    )
    if program_study:
        sp_id = resolve_study_program_id(conn, program_study)
        if sp_id:
            conn.execute("update students set study_program_id = ? where id = ?", (sp_id, student_id))
    return student_id


def _store_import_issue(conn: sqlite3.Connection, issue: dict[str, object], source_file: str) -> None:
    conn.execute(
        """
        insert into import_issues
          (id, sheet_name, row_number, nim, full_name, briva, amount, note, source_file)
        values (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            str(uuid.uuid4()),
            str(issue["sheet_name"]),
            int(issue["row_number"]),
            str(issue.get("nim") or ""),
            str(issue.get("full_name") or ""),
            str(issue.get("briva") or ""),
            str(issue.get("amount") or ""),
            str(issue["note"]),
            source_file,
        ),
    )


def import_workbook(
    workbook_path: str | Path = DEFAULT_WORKBOOK,
    db_path: str | Path = DEFAULT_DB_PATH,
    period: str | None = None,
    bill_type: str = DEFAULT_BILL_TYPE,
    source_file_name: str | None = None,
    confirm_updates: bool = False,
) -> dict[str, object]:
    workbook = Path(workbook_path)
    source_file = source_file_name or workbook.name
    analysis = _analyze_workbook(workbook, db_path, period, source_file)
    if analysis["critical_rows"]:
        raise ValueError("Import dibatalkan karena ada duplikasi atau konflik kritis.")
    if analysis["requires_update_confirmation"] and not confirm_updates:
        raise ValueError("Perubahan nominal atau BRIVA memerlukan konfirmasi admin.")

    conn = connect(db_path)
    init_db(conn)
    issues = 0
    created = 0
    updated = 0
    issue_details: list[dict[str, object]] = []

    with conn:
        conn.execute("delete from import_issues where source_file = ?", (source_file,))
        for issue in analysis["_skipped_issues"]:
            assert isinstance(issue, dict)
            _store_import_issue(conn, issue, source_file)
            issues += 1
            if len(issue_details) < 5:
                issue_details.append({"sheet": issue["sheet_name"], "row_number": issue["row_number"], "note": issue["note"]})
        for action in analysis["actions"]:
            action_type = str(action["type"])
            if action_type == "unchanged":
                continue
            row = action["row"]
            assert isinstance(row, dict)
            nim = str(row["nim"])
            briva = str(row["briva"])
            amount = int(row["amount"])
            row_number = int(row["row_number"])
            due_date = clean_excel_text(row.get("due_date")) or None
            student_id = _upsert_student(conn, row)

            if action_type == "new":
                conn.execute(
                    """
                    insert into bills
                      (id, student_id, briva, amount, period, bill_type, status, instructions, due_date, source_file, source_row_number, updated_at)
                    values (?, ?, ?, ?, ?, ?, 'unpaid', ?, ?, ?, ?, datetime('now'))
                    """,
                    (
                        str(uuid.uuid4()), student_id, briva, amount, str(row["period"]), bill_type,
                        DEFAULT_INSTRUCTIONS, due_date, source_file, row_number,
                    ),
                )
                created += 1
                continue

            existing = action["existing"]
            assert isinstance(existing, sqlite3.Row)
            conn.execute(
                """
                update bills set student_id = ?, briva = ?, amount = ?, period = ?, bill_type = ?, due_date = ?,
                    source_file = ?, source_row_number = ?, updated_at = datetime('now')
                where id = ?
                """,
                (student_id, briva, amount, str(row["period"]), bill_type, due_date, source_file, row_number, existing["id"]),
            )
            updated += 1

        layout = _require_headers(workbook)
        if layout.issue_sheet:
            for record in read_sheet(workbook, layout.issue_sheet):
                issues += 1
                issue = {
                    "sheet_name": layout.issue_sheet,
                    "row_number": int(record.get("_row_number") or 0),
                    "nim": normalize_nim(record.get("NIM")),
                    "full_name": clean_excel_text(record.get("Nama Mahasiswa")),
                    "briva": _normalize_briva(record.get("BRIVA")),
                    "amount": clean_excel_text(record.get("Jumlah")),
                    "note": clean_excel_text(record.get("Keterangan")) or "Data belum lengkap.",
                }
                _store_import_issue(conn, issue, source_file)
                if len(issue_details) < 5:
                    issue_details.append({"sheet": issue["sheet_name"], "row_number": issue["row_number"], "note": issue["note"]})

    conn.close()
    return {
        "imported": created + updated,
        "created": created,
        "updated": updated,
        "unchanged": int(analysis["unchanged_rows"]),
        "issues": issues,
        "issue_details": issue_details,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import data BRIVA UKT ke SQLite.")
    parser.add_argument("--file", default=str(DEFAULT_WORKBOOK), help="Path file .xlsx")
    parser.add_argument("--db", default=str(DEFAULT_DB_PATH), help="Path file SQLite")
    parser.add_argument("--period", help="Opsional. Default mengikuti struktur workbook.")
    parser.add_argument("--confirm-updates", action="store_true", help="Setujui perubahan nominal atau BRIVA.")
    args = parser.parse_args()

    result = import_workbook(args.file, args.db, args.period, confirm_updates=args.confirm_updates)
    print(f"Changed rows: {result['imported']}")
    print(f"Unchanged rows: {result['unchanged']}")
    print(f"Issue rows: {result['issues']}")
    print(f"Database: {Path(args.db).resolve()}")


if __name__ == "__main__":
    main()
